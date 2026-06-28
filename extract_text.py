#!/usr/bin/env python3
"""
Masker 文件文本提取器
用法: python3 extract_text.py <file_path>
输出: 文件的文本内容到 stdout
"""

import sys
import os

def extract_text(filepath: str) -> str:
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".txt":
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    elif ext == ".docx":
        from docx import Document
        doc = Document(filepath)
        paragraphs = []
        for p in doc.paragraphs:
            paragraphs.append(p.text)
        # Also try to extract from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                paragraphs.append(row_text)
        return "\n".join(paragraphs)

    elif ext == ".pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(filepath)
            pages = []
            for page in doc:
                pages.append(page.get_text())
            doc.close()
            return "\n".join(pages)
        except ImportError:
            # Fallback: try pdfminer
            from pdfminer.high_level import extract_text as pdf_extract
            return pdf_extract(filepath)

    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(filepath, data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f"[Sheet: {sheet_name}]")
            for row in ws.iter_rows(values_only=True):
                row_text = "\t".join(str(c) if c is not None else "" for c in row)
                if row_text.strip():
                    lines.append(row_text)
        wb.close()
        return "\n".join(lines)

    elif ext == ".csv":
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    elif ext in (".json", ".xml", ".yaml", ".yml", ".log", ".md"):
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    else:
        raise ValueError(f"Unsupported file type: {ext}")


def main():
    if len(sys.argv) < 2:
        print("Usage: extract_text.py <file_path>", file=sys.stderr)
        sys.exit(1)

    filepath = sys.argv[1]
    if not os.path.isfile(filepath):
        print(f"File not found: {filepath}", file=sys.stderr)
        sys.exit(1)

    try:
        text = extract_text(filepath)
        print(text)
    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        # Still try to read as plain text as last resort
        try:
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                print(f.read())
        except:
            sys.exit(1)


if __name__ == "__main__":
    main()
